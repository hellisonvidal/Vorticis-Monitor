"""Testes da exportação: planilha CSV/XLSX e pacote ZIP de XMLs."""
import csv
import datetime as dt
import zipfile

import openpyxl
import pytest

from monitor_notas.armazenamento import Armazenamento, Documento
from monitor_notas.exportar import planilha, xml_pasta, zip_xml

CNPJ = "12345678000195"
CHAVE_NFE = "35" * 22  # 44 dígitos
CHAVE_NFSE = "5" * 50  # 50 dígitos
CHAVE_RESUMO = "2" * 44


@pytest.fixture
def docs(tmp_path):
    arm = Armazenamento(str(tmp_path / "dados" / "monitor.db"), str(tmp_path / "xml"))
    arm.salvar(Documento(
        origem="nfe", tipo="completo", direcao="recebida", cnpj_empresa=CNPJ,
        chave=CHAVE_NFE, nsu="100", emitente_cnpj="98765432000110",
        emitente_nome="Fornecedor, Ltda", destinatario_cnpj=CNPJ, valor=1234.5,
        data_emissao="2026-10-21T10:30:00-03:00", xml="<nfeProc>completo</nfeProc>",
    ))
    arm.salvar(Documento(
        origem="nfse", tipo="completo", direcao="emitida", cnpj_empresa=CNPJ,
        chave=CHAVE_NFSE, emitente_cnpj=CNPJ, emitente_nome="Nossa Empresa",
        destinatario_cnpj="11222333000181", valor=99.9,
        data_emissao="2026-10-02", xml="<CompNfse>nfse</CompNfse>",
    ))
    arm.salvar(Documento(
        origem="nfe", tipo="evento", direcao="recebida", cnpj_empresa=CNPJ,
        chave=CHAVE_NFE, evento="110111", emitente_cnpj="98765432000110",
        data_emissao="2026-10-22T08:00:00-03:00", xml="<evento>cancelamento</evento>",
    ))
    arm.salvar(Documento(
        origem="nfe", tipo="resumo", direcao="recebida", cnpj_empresa=CNPJ,
        chave=CHAVE_RESUMO, xml="<resNFe/>", data_emissao="data-ruim",
    ))
    return arm.consultar()


CABECALHO = ["Origem", "Direção", "Tipo", "Emissão", "Emitente", "CNPJ emitente",
             "Destinatário CNPJ", "Valor", "Evento", "Chave", "CNPJ empresa", "NSU"]


def _linha_por_chave(linhas, chave, evento=""):
    for l in linhas:
        if l[9] == chave and l[8] == evento:
            return l
    raise AssertionError(f"chave {chave} não encontrada")


def test_csv_separador_virgula_decimal_e_bom(docs, tmp_path):
    caminho = tmp_path / "saida.csv"
    assert planilha(docs, str(caminho)) == 4

    bruto = caminho.read_bytes()
    assert bruto.startswith(b"\xef\xbb\xbf")  # BOM para o Excel

    with open(caminho, encoding="utf-8-sig", newline="") as f:
        linhas = list(csv.reader(f, delimiter=";"))
    assert linhas[0] == CABECALHO
    assert len(linhas) == 5

    nfe = _linha_por_chave(linhas[1:], CHAVE_NFE)
    assert nfe[0] == "NF-e"
    assert nfe[1] == "Recebida"
    assert nfe[2] == "Completo"
    assert nfe[3] == "21/10/2026"
    assert nfe[4] == "Fornecedor, Ltda"
    assert nfe[5] == "98.765.432/0001-10"
    assert nfe[7] == "1234,50"
    assert nfe[10] == "12.345.678/0001-95"
    assert nfe[11] == "100"

    nfse = _linha_por_chave(linhas[1:], CHAVE_NFSE)
    assert nfse[0] == "NFS-e"
    assert nfse[7] == "99,90"


def test_csv_sem_xml(docs, tmp_path):
    caminho = tmp_path / "saida.csv"
    planilha(docs, str(caminho))
    texto = caminho.read_text(encoding="utf-8-sig")
    assert "<nfeProc" not in texto
    assert "xml" not in texto.splitlines()[0].lower()


def test_xlsx_formatos(docs, tmp_path):
    caminho = tmp_path / "saida.xlsx"
    assert planilha(docs, str(caminho)) == 4

    wb = openpyxl.load_workbook(caminho)
    assert wb.sheetnames == ["Documentos"]
    ws = wb["Documentos"]

    assert [c.value for c in ws[1]] == CABECALHO
    assert all(c.font.bold for c in ws[1])
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == f"A1:L{ws.max_row}"
    assert ws.max_row == 5
    assert all(ws.column_dimensions[col].width for col in "ABCDEFGHIJKL")

    linha = next(r for r in range(2, 6) if ws.cell(r, 10).value == CHAVE_NFE and ws.cell(r, 9).value is None)
    valor = ws.cell(linha, 8)
    assert isinstance(valor.value, float)
    assert valor.value == 1234.5
    assert valor.number_format == "#,##0.00"

    data = ws.cell(linha, 4)
    assert isinstance(data.value, dt.datetime)
    assert data.value.date() == dt.date(2026, 10, 21)
    assert data.number_format == "DD/MM/YYYY"

    chave = ws.cell(linha, 10)
    assert chave.value == CHAVE_NFE
    assert chave.data_type == "s"  # texto, não número

    chave_nfse = next(ws.cell(r, 10) for r in range(2, 6) if ws.cell(r, 10).value == CHAVE_NFSE)
    assert chave_nfse.data_type == "s"
    assert len(chave_nfse.value) == 50

    # data inválida é mantida como texto, sem quebrar
    textos = [ws.cell(r, 4).value for r in range(2, 6)]
    assert "data-ruim" in textos


def test_extensao_invalida(docs, tmp_path):
    with pytest.raises(ValueError):
        planilha(docs, str(tmp_path / "saida.txt"))
    with pytest.raises(ValueError):
        planilha(docs, str(tmp_path / "saida"))
    assert not any(tmp_path.glob("saida*"))


def test_zip_nomes_e_conteudo(docs, tmp_path):
    caminho = tmp_path / "xmls.zip"
    assert zip_xml(docs, str(caminho)) == 4

    with zipfile.ZipFile(caminho) as z:
        assert z.testzip() is None
        assert all(i.compress_type == zipfile.ZIP_DEFLATED for i in z.infolist())
        nomes = sorted(z.namelist())
        assert nomes == sorted([
            f"{CNPJ}/nfe/recebida/{CHAVE_NFE}.xml",
            f"{CNPJ}/nfe/recebida/{CHAVE_NFE}-evento110111.xml",
            f"{CNPJ}/nfe/recebida/{CHAVE_RESUMO}-resumo.xml",
            f"{CNPJ}/nfse/emitida/{CHAVE_NFSE}.xml",
        ])
        assert z.read(f"{CNPJ}/nfe/recebida/{CHAVE_NFE}.xml") == b"<nfeProc>completo</nfeProc>"
        assert z.read(f"{CNPJ}/nfe/recebida/{CHAVE_NFE}-evento110111.xml") == b"<evento>cancelamento</evento>"
        assert z.read(f"{CNPJ}/nfse/emitida/{CHAVE_NFSE}.xml") == b"<CompNfse>nfse</CompNfse>"


def test_zip_nomes_duplicados_recebem_sufixo(docs, tmp_path):
    duplicados = list(docs) + list(docs[:1])  # mesmo documento duas vezes
    caminho = tmp_path / "xmls.zip"
    assert zip_xml(duplicados, str(caminho)) == 5
    with zipfile.ZipFile(caminho) as z:
        nomes = z.namelist()
    assert len(nomes) == len(set(nomes)) == 5
    primeiro = docs[0]  # consultar() ordena por data desc: aqui é o resumo
    assert primeiro["chave"] == CHAVE_RESUMO
    assert f"{CNPJ}/nfe/recebida/{CHAVE_RESUMO}-resumo-2.xml" in nomes


def test_xml_pasta_mesma_estrutura_do_zip(docs, tmp_path):
    destino = tmp_path / "saida"
    duplicados = list(docs) + list(docs[:1])
    assert xml_pasta(duplicados, destino) == 5
    arquivos = sorted(p.relative_to(destino).as_posix() for p in destino.rglob("*.xml"))
    assert arquivos == sorted([
        f"{CNPJ}/nfe/recebida/{CHAVE_NFE}.xml",
        f"{CNPJ}/nfe/recebida/{CHAVE_NFE}-evento110111.xml",
        f"{CNPJ}/nfe/recebida/{CHAVE_RESUMO}-resumo.xml",
        f"{CNPJ}/nfe/recebida/{CHAVE_RESUMO}-resumo-2.xml",
        f"{CNPJ}/nfse/emitida/{CHAVE_NFSE}.xml",
    ])
    assert (destino / CNPJ / "nfse" / "emitida" / f"{CHAVE_NFSE}.xml").read_text(
        encoding="utf-8") == "<CompNfse>nfse</CompNfse>"
