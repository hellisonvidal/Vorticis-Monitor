"""Exportação dos documentos armazenados: planilha (CSV ou XLSX) e pacote ZIP de XMLs."""
from __future__ import annotations

import csv
import re
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

CABECALHO = [
    "Origem", "Direção", "Tipo", "Emissão", "Emitente", "CNPJ emitente",
    "Destinatário CNPJ", "Valor", "Evento", "Chave", "CNPJ empresa", "NSU",
]
LARGURAS = [10, 12, 11, 12, 40, 20, 20, 14, 10, 50, 20, 12]
IDX_DATA = 3
IDX_VALOR = 7
IDX_CHAVE = 9
FMT_VALOR = "#,##0.00"
FMT_DATA = "DD/MM/YYYY"

ORIGENS = {"nfe": "NF-e", "nfse": "NFS-e"}
DIRECOES = {"recebida": "Recebida", "emitida": "Emitida"}
TIPOS = {"completo": "Completo", "resumo": "Resumo", "evento": "Evento"}
SUFIXOS = {"completo": "", "resumo": "-resumo"}


def _cnpj(valor) -> str | None:
    if not valor:
        return None
    s = str(valor)
    digitos = re.sub(r"\D", "", s)
    if len(digitos) == 14:
        return f"{digitos[:2]}.{digitos[2:5]}.{digitos[5:8]}/{digitos[8:12]}-{digitos[12:]}"
    return s


def _data(valor) -> date | str | None:
    """Data ISO vira date; valor inválido é mantido como texto; vazio vira None."""
    if not valor:
        return None
    try:
        return date.fromisoformat(str(valor)[:10])
    except ValueError:
        return str(valor)


def _valor(valor) -> float | None:
    if valor is None or valor == "":
        return None
    try:
        return float(valor)
    except (TypeError, ValueError):
        return None


def _linha(r) -> list:
    return [
        ORIGENS.get(r["origem"], r["origem"]),
        DIRECOES.get(r["direcao"], r["direcao"]),
        TIPOS.get(r["tipo"], r["tipo"]),
        _data(r["data_emissao"]),
        r["emitente_nome"] or None,
        _cnpj(r["emitente_cnpj"]),
        _cnpj(r["destinatario_cnpj"]),
        _valor(r["valor"]),
        r["evento"] or None,
        r["chave"],
        _cnpj(r["cnpj_empresa"]),
        r["nsu"] or None,
    ]


def _celula_csv(indice: int, valor) -> str:
    if valor is None:
        return ""
    if indice == IDX_DATA and isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if indice == IDX_VALOR and isinstance(valor, float):
        return f"{valor:.2f}".replace(".", ",")
    return str(valor)


def _escrever_csv(dados: list[list], caminho: str) -> None:
    with open(caminho, "w", encoding="utf-8-sig", newline="") as f:
        escritor = csv.writer(f, delimiter=";")
        escritor.writerow(CABECALHO)
        for linha in dados:
            escritor.writerow([_celula_csv(i, v) for i, v in enumerate(linha)])


def _escrever_xlsx(dados: list[list], caminho: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos"
    ws.append(CABECALHO)
    for celula in ws[1]:
        celula.font = Font(bold=True)
    for linha in dados:
        ws.append(linha)

    ultima = ws.max_row
    for linha in ws.iter_rows(min_row=2, max_row=ultima):
        linha[IDX_VALOR].number_format = FMT_VALOR
        linha[IDX_DATA].number_format = FMT_DATA
        linha[IDX_CHAVE].number_format = "@"  # chave sempre como texto

    for i, largura in enumerate(LARGURAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CABECALHO))}{ultima}"
    wb.save(caminho)


def planilha(linhas, caminho: str) -> int:
    """Grava os documentos em CSV (';', UTF-8 com BOM) ou XLSX. Retorna linhas de dados."""
    minusculo = caminho.lower()
    if not (minusculo.endswith(".csv") or minusculo.endswith(".xlsx")):
        raise ValueError(f"Extensão não suportada (use .csv ou .xlsx): {caminho}")
    dados = [_linha(r) for r in linhas]
    Path(caminho).parent.mkdir(parents=True, exist_ok=True)
    if minusculo.endswith(".csv"):
        _escrever_csv(dados, caminho)
    else:
        _escrever_xlsx(dados, caminho)
    return len(dados)


def _nomes_xml(linhas):
    """Gera (caminho_relativo, linha) únicos: {cnpj}/{origem}/{direcao}/{chave}{sufixo}.xml."""
    usados: set[str] = set()
    for r in linhas:
        sufixo = SUFIXOS.get(r["tipo"], f"-evento{r['evento'] or ''}")
        base = f"{r['cnpj_empresa']}/{r['origem']}/{r['direcao']}/{r['chave']}{sufixo}"
        nome = f"{base}.xml"
        n = 2
        while nome in usados:
            nome = f"{base}-{n}.xml"
            n += 1
        usados.add(nome)
        yield nome, r


def zip_xml(linhas, caminho: str) -> int:
    """Grava um ZIP com um XML por documento em {cnpj}/{origem}/{direcao}/{chave}{sufixo}.xml.
    Retorna o número de arquivos gravados."""
    Path(caminho).parent.mkdir(parents=True, exist_ok=True)
    total = 0
    with zipfile.ZipFile(caminho, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for nome, r in _nomes_xml(linhas):
            z.writestr(nome, r["xml"].encode("utf-8"))
            total += 1
    return total


def xml_pasta(linhas, pasta) -> int:
    """Grava cada XML como arquivo individual em `pasta`, com os mesmos caminhos relativos
    do ZIP. Retorna o número de arquivos gravados."""
    base = Path(pasta)
    total = 0
    for nome, r in _nomes_xml(linhas):
        destino = base / nome
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_text(r["xml"], encoding="utf-8")
        total += 1
    return total
